import unittest
from io import BytesIO

from openpyxl import load_workbook

from backend.services.config_statistics import build_config_statistics
from backend.services.rulebook_export import build_rulebook_pdf, build_rulebook_xlsx


def sample_payload() -> dict:
    model = {
        "checks": {
            "groups": [
                {
                    "id": "address_city",
                    "displayName": "AddressCity (Ort)",
                    "internalName": "address_city",
                    "description": "Validiert den Ort.",
                    "rules": [
                        {
                            "id": "city_empty",
                            "message": "Ort ist leer",
                            "condition": "ADDRESS_CITY IS NULL",
                            "dimension": "Vollständigkeit",
                            "ruleType": "completeness",
                            "active": True,
                        },
                        {
                            "id": "city_short",
                            "message": "Ort ist zu kurz",
                            "condition": "LENGTH(ADDRESS_CITY) < 2",
                            "dimension": "Korrektheit",
                            "ruleType": "min_length",
                            "active": False,
                        },
                    ],
                }
            ]
        }
    }
    statistics = build_config_statistics(
        [
            {
                "report": {
                    "id": "customers",
                    "displayName": "(DEFICIENCIES) Customers",
                    "internalName": "deficiencies_customers",
                },
                "model": model,
                "source": "report",
            }
        ]
    )
    return {
        "configId": "config-1",
        "configName": "Test",
        "configTenant": "test",
        "project": "Master Data",
        "generatedAt": "2026-09-07T10:30:00Z",
        "appVersion": "1.6.1",
        **statistics,
    }


class RulebookExportTest(unittest.TestCase):
    def test_statistics_contains_complete_group_and_rule_details(self) -> None:
        payload = sample_payload()

        group = payload["reports"][0]["groups"][0]
        self.assertEqual("AddressCity (Ort)", group["displayName"])
        self.assertEqual(2, group["ruleCount"])
        self.assertEqual("Ort ist leer", group["rules"][0]["message"])
        self.assertEqual("completeness", group["rules"][0]["ruleType"])
        self.assertEqual([{"name": "Korrektheit", "count": 1}, {"name": "Vollständigkeit", "count": 1}], payload["dimensions"])

    def test_pdf_contains_report_group_and_rules(self) -> None:
        content = build_rulebook_pdf(sample_payload(), "de")

        self.assertTrue(content.startswith(b"%PDF"))
        self.assertGreater(len(content), 2_000)

    def test_excel_contains_filterable_rule_details_and_summary(self) -> None:
        content = build_rulebook_xlsx(sample_payload(), "de")

        workbook = load_workbook(BytesIO(content))
        self.assertEqual(["Übersicht", "Regeln"], workbook.sheetnames)
        details = workbook["Regeln"]
        self.assertEqual("Berichte", details["A1"].value)
        self.assertEqual("(DEFICIENCIES) Customers", details["A2"].value)
        self.assertEqual("AddressCity (Ort)", details["C2"].value)
        self.assertEqual("Aktiv", details["G2"].value)
        self.assertTrue(details.tables)
        self.assertEqual("A2", details.freeze_panes)


if __name__ == "__main__":
    unittest.main()
