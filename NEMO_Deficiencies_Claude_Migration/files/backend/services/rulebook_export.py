from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table as ExcelTable, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    CondPageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table as PdfTable,
    TableStyle,
)
from xml.sax.saxutils import escape


RULE_TYPE_LABELS = {
    "completeness": {"de": "Pflichtfeld / Leerprüfung", "en": "Required field / empty check"},
    "trim_whitespace": {"de": "Leerzeichen bereinigen", "en": "Trim whitespace"},
    "min_length": {"de": "Mindestlänge", "en": "Minimum length"},
    "max_length": {"de": "Maximallänge", "en": "Maximum length"},
    "regex": {"de": "Regex-Regel", "en": "Regex rule"},
    "obsolete_terms": {"de": "Obsolete Begriffe", "en": "Obsolete terms"},
    "mixed_umlaut_spelling": {"de": "Umlaut-Schreibweise", "en": "Umlaut spelling"},
    "legal_form_normalization": {"de": "Rechtsform-Normalisierung", "en": "Legal form normalization"},
    "custom": {"de": "Sonderregel", "en": "Custom rule"},
}

COLORS = {
    "navy": "243047",
    "blue": "4F7DF3",
    "light_blue": "EAF0FF",
    "green": "2F855A",
    "light_green": "E8F5EE",
    "red": "C2414A",
    "light_red": "FDECEC",
    "gray": "667085",
    "light_gray": "F3F5F8",
    "border": "D5DAE3",
    "white": "FFFFFF",
}


def _language(value: str) -> str:
    return "en" if str(value).lower() == "en" else "de"


def _labels(language: str) -> Dict[str, str]:
    if _language(language) == "en":
        return {
            "title": "NEMO data quality rulebook",
            "configuration": "Configuration",
            "tenant": "Tenant",
            "project": "Project",
            "created": "Created",
            "version": "Application version",
            "reports": "Reports",
            "groups": "Rule groups",
            "rules": "Rules",
            "active": "Active",
            "inactive": "Inactive",
            "field": "Rule group / field",
            "message": "Rule / error message",
            "status": "Status",
            "dimension": "DQ type",
            "rule_type": "Rule type",
            "condition": "Condition",
            "source": "Source",
            "draft": "Draft",
            "nemo": "NEMO",
            "summary": "Overview",
            "details": "Rules",
            "skipped": "Skipped reports",
            "no_rules": "No rules",
            "distribution": "Rule status",
        }
    return {
        "title": "NEMO-Datenqualitätsregelwerk",
        "configuration": "Konfiguration",
        "tenant": "Tenant",
        "project": "Projekt",
        "created": "Erstellt",
        "version": "Anwendungsversion",
        "reports": "Berichte",
        "groups": "Regelgruppen",
        "rules": "Regeln",
        "active": "Aktiv",
        "inactive": "Inaktiv",
        "field": "Regelgruppe / Feld",
        "message": "Regel / Fehlermeldung",
        "status": "Status",
        "dimension": "DQ-Typ",
        "rule_type": "Regeltyp",
        "condition": "Bedingung",
        "source": "Quelle",
        "draft": "Entwurf",
        "nemo": "NEMO",
        "summary": "Übersicht",
        "details": "Regeln",
        "skipped": "Übersprungene Berichte",
        "no_rules": "Keine Regeln",
        "distribution": "Regelstatus",
    }


def _rule_type_label(rule_type: str, language: str) -> str:
    labels = RULE_TYPE_LABELS.get(rule_type)
    if not labels:
        return rule_type or RULE_TYPE_LABELS["custom"][_language(language)]
    return labels[_language(language)]


def _created_text(value: str) -> str:
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.strftime("%d.%m.%Y %H:%M")
    except (TypeError, ValueError):
        return str(value or "")


def _iter_rules(payload: Dict[str, Any]) -> Iterable[Dict[str, Any]]:
    for report in payload.get("reports") or []:
        for group in report.get("groups") or []:
            rules = group.get("rules") or []
            if not rules:
                yield {"report": report, "group": group, "rule": None}
            for rule in rules:
                yield {"report": report, "group": group, "rule": rule}


def build_rulebook_pdf(payload: Dict[str, Any], language: str = "de") -> bytes:
    language = _language(language)
    labels = _labels(language)
    buffer = BytesIO()
    page_width, _page_height = landscape(A4)
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=13 * mm,
        title=labels["title"],
        author="NEMO Deficiencies",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "RulebookTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=19,
        leading=22,
        textColor=colors.HexColor(f"#{COLORS['navy']}"),
        alignment=TA_LEFT,
        spaceAfter=5 * mm,
    )
    report_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor(f"#{COLORS['navy']}"),
        spaceBefore=4 * mm,
        spaceAfter=2 * mm,
    )
    group_style = ParagraphStyle(
        "Group",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor(f"#{COLORS['navy']}"),
    )
    cell_style = ParagraphStyle(
        "Cell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=7.2,
        leading=8.6,
        textColor=colors.HexColor("#202735"),
    )
    small_style = ParagraphStyle(
        "Small",
        parent=cell_style,
        fontSize=6.3,
        leading=7.6,
        textColor=colors.HexColor(f"#{COLORS['gray']}"),
    )

    story = [Paragraph(escape(labels["title"]), title_style)]
    metadata = [
        [labels["configuration"], payload.get("configName") or payload.get("configId") or ""],
        [labels["tenant"], payload.get("configTenant") or ""],
        [labels["project"], payload.get("project") or ""],
        [labels["created"], _created_text(payload.get("generatedAt") or "")],
        [labels["version"], payload.get("appVersion") or ""],
    ]
    meta_table = PdfTable(metadata, colWidths=[34 * mm, 105 * mm], hAlign="LEFT")
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor(f"#{COLORS['gray']}")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.extend([meta_table, Spacer(1, 5 * mm)])

    summary = payload.get("summary") or {}
    metric_values = [
        (labels["reports"], summary.get("reportCount", 0)),
        (labels["groups"], summary.get("groupCount", 0)),
        (labels["rules"], summary.get("ruleCount", 0)),
        (labels["active"], summary.get("activeRuleCount", 0)),
        (labels["inactive"], summary.get("inactiveRuleCount", 0)),
    ]
    metric_table = PdfTable(
        [[Paragraph(escape(label), small_style) for label, _value in metric_values],
         [str(value) for _label, value in metric_values]],
        colWidths=[34 * mm] * len(metric_values),
        hAlign="LEFT",
    )
    metric_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{COLORS['light_gray']}")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{COLORS['border']}")),
                ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{COLORS['border']}")),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 1), (-1, 1), 14),
                ("TEXTCOLOR", (3, 1), (3, 1), colors.HexColor(f"#{COLORS['green']}")),
                ("TEXTCOLOR", (4, 1), (4, 1), colors.HexColor(f"#{COLORS['red']}")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([metric_table, Spacer(1, 4 * mm)])

    for report in payload.get("reports") or []:
        story.append(CondPageBreak(38 * mm))
        source_label = labels["draft"] if report.get("source") == "draft" else labels["nemo"]
        report_title = (
            f"{escape(str(report.get('displayName') or report.get('internalName') or ''))} "
            f"<font color='#{COLORS['gray']}' size='7'>"
            f"{escape(str(report.get('internalName') or ''))} | {source_label}</font>"
        )
        story.append(Paragraph(report_title, report_style))
        groups = report.get("groups") or []
        if not groups:
            story.append(Paragraph(labels["no_rules"], cell_style))
            continue
        for group in groups:
            story.append(CondPageBreak(24 * mm))
            group_title = escape(str(group.get("displayName") or group.get("internalName") or ""))
            internal_name = escape(str(group.get("internalName") or ""))
            description = escape(str(group.get("description") or ""))
            group_text = f"{group_title} <font color='#{COLORS['gray']}'>({internal_name})</font>"
            if description:
                group_text += f"<br/><font name='Helvetica' color='#{COLORS['gray']}'>{description}</font>"
            story.append(Paragraph(group_text, group_style))
            header = [labels["message"], labels["status"], labels["dimension"], labels["rule_type"], labels["condition"]]
            rows = [header]
            for rule in group.get("rules") or []:
                rows.append(
                    [
                        Paragraph(escape(str(rule.get("message") or "")), cell_style),
                        labels["active"] if rule.get("active") else labels["inactive"],
                        Paragraph(escape(str(rule.get("dimension") or "")), cell_style),
                        Paragraph(escape(_rule_type_label(str(rule.get("ruleType") or ""), language)), cell_style),
                        Paragraph(escape(str(rule.get("condition") or "")), small_style),
                    ]
                )
            if len(rows) == 1:
                rows.append([labels["no_rules"], "", "", "", ""])
            table = PdfTable(
                rows,
                colWidths=[54 * mm, 18 * mm, 27 * mm, 38 * mm, page_width - 24 * mm - 137 * mm],
                repeatRows=1,
                hAlign="LEFT",
            )
            style_commands = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{COLORS['navy']}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 7),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor(f"#{COLORS['border']}")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
            for row_index, rule in enumerate(group.get("rules") or [], start=1):
                background = COLORS["light_green"] if rule.get("active") else COLORS["light_red"]
                status_color = COLORS["green"] if rule.get("active") else COLORS["red"]
                style_commands.extend(
                    [
                        ("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor(f"#{background}")),
                        ("TEXTCOLOR", (1, row_index), (1, row_index), colors.HexColor(f"#{status_color}")),
                        ("FONTNAME", (1, row_index), (1, row_index), "Helvetica-Bold"),
                    ]
                )
            table.setStyle(TableStyle(style_commands))
            story.extend([table, Spacer(1, 2.5 * mm)])

    skipped = payload.get("skippedReports") or []
    if skipped:
        story.append(CondPageBreak(25 * mm))
        story.append(Paragraph(labels["skipped"], report_style))
        for item in skipped:
            story.append(
                Paragraph(
                    f"{escape(str(item.get('reportRef') or ''))}: {escape(str(item.get('error') or ''))}",
                    small_style,
                )
            )

    def add_page_number(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor(f"#{COLORS['gray']}"))
        canvas.drawString(12 * mm, 7 * mm, "NEMO Deficiencies")
        canvas.drawRightString(page_width - 12 * mm, 7 * mm, f"{document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
    return buffer.getvalue()


def build_rulebook_xlsx(payload: Dict[str, Any], language: str = "de") -> bytes:
    language = _language(language)
    labels = _labels(language)
    workbook = Workbook()
    overview = workbook.active
    overview.title = labels["summary"]
    details = workbook.create_sheet(labels["details"])

    overview.sheet_view.showGridLines = False
    overview["A1"] = labels["title"]
    overview["A1"].font = Font(size=20, bold=True, color=COLORS["navy"])
    overview.merge_cells("A1:F1")
    metadata = [
        (labels["configuration"], payload.get("configName") or payload.get("configId") or ""),
        (labels["tenant"], payload.get("configTenant") or ""),
        (labels["project"], payload.get("project") or ""),
        (labels["created"], _created_text(payload.get("generatedAt") or "")),
        (labels["version"], payload.get("appVersion") or ""),
    ]
    for row, (label, value) in enumerate(metadata, start=3):
        overview.cell(row, 1, label).font = Font(bold=True, color=COLORS["gray"])
        overview.cell(row, 2, value)

    summary = payload.get("summary") or {}
    metrics = [
        (labels["reports"], summary.get("reportCount", 0)),
        (labels["groups"], summary.get("groupCount", 0)),
        (labels["rules"], summary.get("ruleCount", 0)),
        (labels["active"], summary.get("activeRuleCount", 0)),
        (labels["inactive"], summary.get("inactiveRuleCount", 0)),
    ]
    for column, (label, value) in enumerate(metrics, start=1):
        header_cell = overview.cell(9, column, label)
        header_cell.font = Font(bold=True, color=COLORS["white"])
        header_cell.fill = PatternFill("solid", fgColor=COLORS["navy"])
        header_cell.alignment = Alignment(horizontal="center")
        value_cell = overview.cell(10, column, value)
        value_cell.font = Font(size=16, bold=True, color=COLORS["green"] if column == 4 else COLORS["red"] if column == 5 else COLORS["navy"])
        value_cell.alignment = Alignment(horizontal="center")

    report_headers = [labels["reports"], "Internalname", labels["groups"], labels["rules"], labels["active"], labels["inactive"], labels["source"]]
    report_start = 13
    for column, header in enumerate(report_headers, start=1):
        overview.cell(report_start, column, header)
    for row, report in enumerate(payload.get("reports") or [], start=report_start + 1):
        values = [
            report.get("displayName") or "",
            report.get("internalName") or "",
            report.get("groupCount") or 0,
            report.get("ruleCount") or 0,
            report.get("activeRuleCount") or 0,
            report.get("inactiveRuleCount") or 0,
            labels["draft"] if report.get("source") == "draft" else labels["nemo"],
        ]
        for column, value in enumerate(values, start=1):
            overview.cell(row, column, value)
    report_end = report_start + max(1, len(payload.get("reports") or []))
    if payload.get("reports"):
        table = ExcelTable(displayName="ReportSummary", ref=f"A{report_start}:G{report_end}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        overview.add_table(table)

    overview["I9"] = labels["status"]
    overview["J9"] = labels["rules"]
    overview["I10"] = labels["active"]
    overview["J10"] = summary.get("activeRuleCount", 0)
    overview["I11"] = labels["inactive"]
    overview["J11"] = summary.get("inactiveRuleCount", 0)
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = labels["distribution"]
    chart.height = 5.2
    chart.width = 9.2
    chart.legend = None
    chart.add_data(Reference(overview, min_col=10, min_row=9, max_row=11), titles_from_data=True)
    chart.set_categories(Reference(overview, min_col=9, min_row=10, max_row=11))
    overview.add_chart(chart, "I13")

    detail_headers = [
        labels["reports"],
        "Report Internalname",
        labels["field"],
        "Feld Internalname",
        "Beschreibung" if language == "de" else "Description",
        labels["message"],
        labels["status"],
        labels["dimension"],
        labels["rule_type"],
        "Regeltyp ID" if language == "de" else "Rule type ID",
        labels["condition"],
        labels["source"],
    ]
    details.append(detail_headers)
    for item in _iter_rules(payload):
        report = item["report"]
        group = item["group"]
        rule = item["rule"] or {}
        details.append(
            [
                report.get("displayName") or "",
                report.get("internalName") or "",
                group.get("displayName") or "",
                group.get("internalName") or "",
                group.get("description") or "",
                rule.get("message") or labels["no_rules"],
                (labels["active"] if rule.get("active") else labels["inactive"]) if item["rule"] else "",
                rule.get("dimension") or "",
                _rule_type_label(str(rule.get("ruleType") or ""), language) if item["rule"] else "",
                rule.get("ruleType") or "",
                rule.get("condition") or "",
                labels["draft"] if report.get("source") == "draft" else labels["nemo"],
            ]
        )

    header_fill = PatternFill("solid", fgColor=COLORS["navy"])
    header_font = Font(bold=True, color=COLORS["white"])
    thin_border = Border(bottom=Side(style="thin", color=COLORS["border"]))
    for cell in details[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in details.iter_rows(min_row=2):
        row[6].fill = PatternFill("solid", fgColor=COLORS["light_green"] if row[6].value == labels["active"] else COLORS["light_red"] if row[6].value else COLORS["white"])
        row[6].font = Font(bold=True, color=COLORS["green"] if row[6].value == labels["active"] else COLORS["red"])
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
    if details.max_row > 1:
        detail_table = ExcelTable(displayName="RuleDetails", ref=f"A1:L{details.max_row}")
        detail_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        details.add_table(detail_table)
    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:L{details.max_row}"
    details.sheet_view.showGridLines = False

    overview.freeze_panes = f"A{report_start + 1}"
    overview.column_dimensions["A"].width = 38
    overview.column_dimensions["B"].width = 34
    for column in "CDEFG":
        overview.column_dimensions[column].width = 15
    widths = [34, 30, 32, 26, 42, 38, 12, 20, 28, 20, 70, 14]
    for index, width in enumerate(widths, start=1):
        details.column_dimensions[chr(64 + index)].width = width
    details.row_dimensions[1].height = 32
    details.sheet_properties.pageSetUpPr.fitToPage = True
    details.page_setup.orientation = "landscape"
    details.page_setup.fitToWidth = 1
    details.page_setup.fitToHeight = 0
    details.print_title_rows = "1:1"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
